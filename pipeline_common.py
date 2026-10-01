"""Shared Excel pipeline for main_parallel.py / step1only_main_parallel.py / step2only_main_parallel.py."""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import re
import sys
import traceback
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from tqdm import tqdm

from copilot_common import (
    DEFAULT_MAX_RETRIES,
    DEFAULT_REQUEST_TIMEOUT_SEC,
    CopilotRunner,
    build_prompt,
    interpret_screening,
    read_prompt_text,
    resolve_token,
    unwrap_extraction,
    validate_extraction,
    validate_screening,
)

EXIT_OK = 0
EXIT_FATAL = 1
EXIT_ROW_ERRORS = 2

DEFAULT_MODEL = "auto"
DEFAULT_INPUT_COLUMNS = ["請求項（英語）", "タイトル（英語）"]
DEFAULT_SHEET = "savedrecs"
SAVE_EVERY_ROWS = 10
DEFAULT_ROW_RETRY_PASSES = 2
DEFAULT_ROW_RETRY_DELAY_SEC = 5.0
EXCEL_CELL_MAX_CHARS = 32767
METADATA_CHUNK_CHARS = 30000
METADATA_SHEET = "実行条件"

# stdout markers parsed by streamlit_app.py
OUTPUT_MARKER = "[OUTPUT]"
MERGE_DONE_MARKER = "[MERGE DONE]"


class Abort(Exception):
    pass


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args(description: str) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=description)
    p.add_argument("--input", type=str, help="入力Excelのパス（指定時はUIスキップ）")
    p.add_argument("--sheet", type=str, help="シート名（指定時はUIスキップ）")
    p.add_argument("--prompt1", type=str, help="1段階目（スクリーニング）プロンプトTXTのパス")
    p.add_argument("--prompt2", type=str, help="2段階目（詳細抽出）プロンプトTXTのパス")
    p.add_argument("--apikey", "--token", dest="apikey", type=str, default="",
                   help="Copilot用GitHubトークンのファイルパス（省略時は環境変数、なければログイン済みユーザー）")
    p.add_argument("--outdir", type=str, help="出力先フォルダ（省略時は入力Excelと同じ場所）")
    p.add_argument("--no-ui", dest="no_ui", action="store_true", help="UIダイアログを出さない（完全ヘッドレス）")
    p.add_argument("--model-stage1", type=str, default=DEFAULT_MODEL, help="1段階目で使うCopilotモデル名")
    p.add_argument("--model-stage2", type=str, default=DEFAULT_MODEL, help="2段階目で使うCopilotモデル名")
    p.add_argument("--input-column", dest="input_columns", action="append",
                   help="プロンプトの [入力N] に差し込むExcel列名。複数指定可")
    p.add_argument("--web-search", dest="web_search", action="store_true",
                   help="互換オプション（Copilot SDKでは無視されます）")
    p.add_argument("--max-retries", type=int, default=DEFAULT_MAX_RETRIES, help="1リクエストあたりの最大試行回数")
    p.add_argument("--request-timeout", type=float, default=DEFAULT_REQUEST_TIMEOUT_SEC,
                   help="1リクエストあたりのタイムアウト秒")
    p.add_argument("--row-retry-passes", type=int, default=DEFAULT_ROW_RETRY_PASSES,
                   help="エラーになった行だけを対象に、全行処理後にもう一度やり直す回数")
    p.add_argument("--row-retry-delay", type=float, default=DEFAULT_ROW_RETRY_DELAY_SEC,
                   help="やり直しパスの前に空ける秒数")
    p.add_argument("--cli-path", type=str, help="Copilot CLI の実行ファイル（省略時はSDK同梱版）")
    p.add_argument("--cli-url", type=str, help="起動済み Copilot CLI サーバーのURL")
    p.add_argument("--merge-after", dest="merge_after", action="store_true",
                   help="この実行の最後に Part1〜PartN の最新出力を1つのExcelに結合する")
    p.add_argument("--parts", type=int, default=10, help="パート数（既定10）")
    p.add_argument("--merge-only", dest="merge_only", action="store_true",
                   help="処理は行わず、Part1〜PartN の最新出力を結合して終了する")
    p.add_argument("--merge-base", type=str,
                   help="結合のベース名（拡張子なし・末尾が PartN）。例: 標的_…_Part10")
    p.add_argument("--run-metadata", type=str,
                   help="MERGEDファイルの実行条件シートに記録するJSONファイルのパス")
    return p.parse_args()


class Notifier:
    """Shows tkinter dialogs in UI mode, otherwise prints."""

    def __init__(self, use_ui: bool) -> None:
        self.use_ui = use_ui
        self._tk: Any = None
        self._root: Any = None
        if use_ui:
            try:
                import tkinter as tk
                from tkinter import filedialog, messagebox, simpledialog

                self._tk = (filedialog, messagebox, simpledialog)
                self._root = tk.Tk()
                self._root.withdraw()
            except Exception:
                print("[WARN] tkinter が利用できないため --no-ui モードで実行します", file=sys.stderr)
                self.use_ui = False

    def info(self, title: str, msg: str) -> None:
        if self.use_ui:
            try:
                self._tk[1].showinfo(title, msg)
                return
            except Exception:
                pass
        print(f"[INFO] {title}: {msg}")

    def error(self, title: str, msg: str) -> None:
        if self.use_ui:
            try:
                self._tk[1].showerror(title, msg)
                return
            except Exception:
                pass
        print(f"[ERROR] {title}: {msg}", file=sys.stderr)

    def select_file(self, title: str, filetypes: list[tuple[str, str]]) -> str:
        if not self.use_ui:
            return ""
        return self._tk[0].askopenfilename(title=title, filetypes=filetypes) or ""

    def ask_string(self, title: str, prompt: str) -> str:
        if not self.use_ui:
            return ""
        return self._tk[2].askstring(title, prompt, parent=self._root) or ""


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(text) > EXCEL_CELL_MAX_CHARS:
        suffix = "…[truncated]"
        text = text[: EXCEL_CELL_MAX_CHARS - len(suffix)] + suffix
    return text


def cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def rightmost_filled_col(ws, header_row: int = 1) -> int:
    for col in range(ws.max_column, 0, -1):
        val = ws.cell(row=header_row, column=col).value
        if val is not None and str(val).strip() != "":
            return col
    return 0


def header_columns(ws, header_row: int = 1) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        name = "" if val is None else str(val).strip()
        if name and name not in mapping:
            mapping[name] = col
    return mapping


def flatten_record(data: dict[str, Any], parent: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in data.items():
        name = f"{parent}.{key}" if parent else str(key)
        if isinstance(value, dict):
            out.update(flatten_record(value, name))
        elif isinstance(value, list):
            if all(not isinstance(v, (dict, list)) for v in value):
                out[name] = "; ".join("" if v is None else str(v) for v in value)
            else:
                out[name] = json.dumps(value, ensure_ascii=False)
        else:
            out[name] = value
    return out


class DynamicColumns:
    """Appends new header columns on demand, starting at start_col."""

    def __init__(self, ws, start_col: int, prefix: str, header_row: int = 1) -> None:
        self.ws = ws
        self.prefix = prefix
        self.header_row = header_row
        self.next_col = start_col
        self.columns: dict[str, int] = {}

    def write(self, row: int, record: dict[str, Any]) -> None:
        for key, value in record.items():
            name = f"{self.prefix}{key}"
            col = self.columns.get(name)
            if col is None:
                col = self.next_col
                self.next_col += 1
                self.columns[name] = col
                self.ws.cell(row=self.header_row, column=col, value=excel_safe(name))
            self.ws.cell(row=row, column=col, value=excel_safe(value))


def file_sha1(path: str) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# Merge
# ---------------------------------------------------------------------------

_PART_TAIL = re.compile(r"^(.*?Part)(\d+)$")
_TIMESTAMP_PREFIX = r"\d{4}-\d{2}-\d{2}_\d{6}_"


def split_base_and_part(basename_noext: str) -> tuple[str | None, int | None]:
    m = _PART_TAIL.search(basename_noext)
    if not m:
        return None, None
    return m.group(1), int(m.group(2))


def find_latest_part_output(outdir: str, target_base: str) -> str | None:
    """Latest 'YYYY-mm-dd_HHMMSS_<target_base>.xlsx' in outdir (exact base match)."""
    pattern = re.compile(rf"^{_TIMESTAMP_PREFIX}{re.escape(target_base)}\.xlsx$")
    candidates = [
        os.path.join(outdir, name)
        for name in os.listdir(outdir)
        if pattern.match(name)
    ]
    if not candidates:
        return None
    return max(candidates, key=os.path.getmtime)


def _load_run_metadata(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {"metadata": data}
    except Exception as exc:
        return {"metadata_load_error": f"{path}: {exc}"}


def _append_metadata_value(ws, key: str, value: Any) -> None:
    if value is None:
        text = ""
    elif isinstance(value, (list, dict)):
        text = json.dumps(value, ensure_ascii=False, indent=2)
    else:
        text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if len(text) <= METADATA_CHUNK_CHARS:
        ws.append([key, text])
        return
    total = (len(text) + METADATA_CHUNK_CHARS - 1) // METADATA_CHUNK_CHARS
    for idx in range(total):
        chunk = text[idx * METADATA_CHUNK_CHARS:(idx + 1) * METADATA_CHUNK_CHARS]
        ws.append([f"{key} ({idx + 1}/{total})", chunk])


def _append_run_metadata_sheet(wb_out, metadata: dict[str, Any], data_sheet: str) -> None:
    title = METADATA_SHEET if METADATA_SHEET != data_sheet else f"{METADATA_SHEET}_meta"
    ws_meta = wb_out.create_sheet(title=title)
    ws_meta.append(["項目", "値"])
    ordered_keys = [
        "generated_at", "merged_at", "mode", "parts", "base_excel_path", "base_name",
        "sheet_name", "input_columns", "prompt1_path", "prompt1_sha1", "prompt1_text",
        "prompt2_path", "prompt2_sha1", "prompt2_text", "auth", "out_dir",
        "model_screening", "model_extraction", "python_command",
    ]
    written: set[str] = set()
    for key in ordered_keys:
        if key in metadata:
            _append_metadata_value(ws_meta, key, metadata[key])
            written.add(key)
    for key in sorted(k for k in metadata if k not in written):
        _append_metadata_value(ws_meta, key, metadata[key])


def merge_parts_into_single(
    outdir: str,
    input_base_noext: str,
    sheetname: str,
    parts: int,
    run_metadata_path: str | None = None,
) -> str:
    """
    Merge the latest output of Part1..PartN into one workbook.
    Headers are unioned across parts (first-seen order); missing parts are an error.
    """
    prefix_part, _ = split_base_and_part(input_base_noext)
    if prefix_part is None:
        raise ValueError(f"入力ベース名に 'PartN' が見つかりません: {input_base_noext}")

    part_files: list[tuple[int, str]] = []
    missing: list[int] = []
    for part_no in range(1, parts + 1):
        path = find_latest_part_output(outdir, f"{prefix_part}{part_no}")
        if path is None:
            missing.append(part_no)
        else:
            part_files.append((part_no, path))
    if missing:
        raise FileNotFoundError(
            f"以下のPart出力が見つかりませんでした: {missing}\n"
            f"フォルダ: {outdir}\n"
            f"期待パターン: 'YYYY-mm-dd_HHMMSS_{prefix_part}N.xlsx'"
        )

    union_headers: list[str] = []
    seen: set[str] = set()
    loaded: list[tuple[list[str], list[tuple[Any, ...]]]] = []
    for part_no, path in part_files:
        print(f"[MERGE] Part{part_no}: {path}")
        wb_in = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheetname not in wb_in.sheetnames:
                raise ValueError(f"シート '{sheetname}' が見つかりません: {path}")
            rows = wb_in[sheetname].iter_rows(values_only=True)
            header_row = next(rows, ())
            headers = ["" if v is None else str(v).strip() for v in header_row]
            data_rows = list(rows)
        finally:
            wb_in.close()
        for h in headers:
            if h and h not in seen:
                seen.add(h)
                union_headers.append(h)
        loaded.append((headers, data_rows))

    if not union_headers:
        raise ValueError("結合対象ファイルからヘッダーを取得できませんでした。")

    base_prefix = re.sub(r"Part$", "", prefix_part)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    merged_path = os.path.join(outdir, f"{timestamp}_{base_prefix}MERGED.xlsx")

    wb_out = Workbook(write_only=True)
    ws_out = wb_out.create_sheet(title=sheetname)
    ws_out.append(union_headers)
    union_index = {h: i for i, h in enumerate(union_headers)}
    for headers, data_rows in loaded:
        mapping = {src: union_index[h] for src, h in enumerate(headers) if h in union_index}
        for row in data_rows:
            out_row: list[Any] = [None] * len(union_headers)
            for src, value in enumerate(row):
                dst = mapping.get(src)
                if dst is not None:
                    out_row[dst] = value
            ws_out.append(out_row)

    metadata = _load_run_metadata(run_metadata_path)
    metadata.setdefault("merged_at", datetime.now().isoformat(timespec="seconds"))
    metadata.setdefault("merge_base", input_base_noext)
    metadata.setdefault("merged_output_path", merged_path)
    _append_run_metadata_sheet(wb_out, metadata, sheetname)

    wb_out.save(merged_path)
    wb_out.close()
    return merged_path


def _merge_only(args: argparse.Namespace, ui: Notifier) -> int:
    if not args.outdir:
        ui.error("結合エラー", "--outdir を指定してください")
        return EXIT_FATAL
    if not args.merge_base:
        ui.error("結合エラー", "--merge-base を指定してください（例: 標的_…_Part10）")
        return EXIT_FATAL
    try:
        merged_path = merge_parts_into_single(
            outdir=args.outdir,
            input_base_noext=args.merge_base,
            sheetname=args.sheet or "Sheet1",
            parts=args.parts,
            run_metadata_path=args.run_metadata,
        )
    except Exception as exc:
        ui.error("結合エラー", str(exc))
        return EXIT_FATAL
    print(f"{MERGE_DONE_MARKER} {merged_path}")
    ui.info("結合完了", f"Part1〜Part{args.parts} を結合しました:\n{merged_path}")
    return EXIT_OK


# ---------------------------------------------------------------------------
# Row processing
# ---------------------------------------------------------------------------

@dataclass
class JobConfig:
    input_path: str
    sheetname: str
    prompt1: str | None
    prompt2: str | None
    token: str
    outdir: str
    input_columns: list[str]


def _resolve_job(args: argparse.Namespace, ui: Notifier, run_stage1: bool, run_stage2: bool) -> JobConfig:
    input_path = args.input
    if not input_path:
        if not ui.use_ui:
            raise Abort("--input を指定するか、UI有効で実行してください。")
        ui.info("ToDo", "処理対象のファイル (Excel) を選択してください。")
        input_path = ui.select_file("ファイルを選択", [("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
    if not input_path:
        raise Abort("ファイルが選択されませんでした。")
    if not os.path.isfile(input_path):
        raise Abort(f"元ファイルが見つかりません: {input_path}")

    sheetname = args.sheet
    if not sheetname:
        sheetname = ui.ask_string("シート名入力", "記入用シートの名前を入力してください:") or DEFAULT_SHEET

    def pick_prompt(given: str | None, label: str) -> str:
        path = given
        if not path:
            if not ui.use_ui:
                raise Abort(f"{label} のプロンプトを指定するか、UI有効で実行してください。")
            ui.info("ToDo", f"【{label}】プロンプトファイル(.txt)を選択してください。")
            path = ui.select_file(f"{label}プロンプトを選択", [("Text files", "*.txt"), ("All files", "*.*")])
        if not path:
            raise Abort(f"{label}のプロンプトが選択されませんでした。")
        path = os.path.abspath(path)
        if not os.path.isfile(path):
            raise Abort(f"{label}プロンプトが見つかりません: {path}")
        return path

    prompt1 = pick_prompt(args.prompt1, "1段階目（スクリーニング）") if run_stage1 else None
    prompt2 = pick_prompt(args.prompt2, "2段階目（詳細抽出）") if run_stage2 else None

    try:
        token = resolve_token(args.apikey)
    except FileNotFoundError as exc:
        raise Abort(str(exc)) from exc

    outdir = args.outdir or os.path.dirname(os.path.abspath(input_path))
    os.makedirs(outdir, exist_ok=True)

    cols = [c.strip() for c in (args.input_columns or []) if c and c.strip()]
    return JobConfig(
        input_path=input_path,
        sheetname=sheetname,
        prompt1=prompt1,
        prompt2=prompt2,
        token=token,
        outdir=outdir,
        input_columns=cols or list(DEFAULT_INPUT_COLUMNS),
    )


def _save(wb, path: str, final: bool = False) -> bool:
    try:
        wb.save(path)
        return True
    except Exception as exc:
        level = "ERROR" if final else "WARN"
        print(f"[{level}] Excel保存失敗: {path}: {exc}", file=sys.stderr)
        return False


async def _process_row(
    runner: CopilotRunner,
    ws,
    r: int,
    values: list[str],
    template1: str,
    template2: str,
    args: argparse.Namespace,
    run_stage1: bool,
    run_stage2: bool,
    cols: tuple[int, int, int, int],
    stage2_cols: "DynamicColumns | None",
) -> bool:
    """Process one row in place. Returns True if stage1 or stage2 failed for this row."""
    s1_judge, s1_reason, s1_raw, s2_status = cols
    preview = next((v for v in reversed(values[:2]) if v.strip()), "")[:50]
    print(f"\n--- 処理中の行: {r} | 入力プレビュー: {preview}... ---", flush=True)

    relevant = True
    row_error = False
    stage1_failed = False
    if run_stage1:
        res1 = await runner.generate_json(build_prompt(template1, values), args.model_stage1, validate_screening)
        if res1.ok:
            relevant, reason = interpret_screening(res1.data)
            ws.cell(row=r, column=s1_judge, value=str(relevant))
            ws.cell(row=r, column=s1_reason, value=excel_safe(reason))
            ws.cell(row=r, column=s1_raw).value = None  # clear a stale error from a previous retry pass
            print(f"[SCREENING] row={r} result={relevant}")
        else:
            row_error = True
            relevant = False
            stage1_failed = True
            ws.cell(row=r, column=s1_judge, value="ERROR")
            ws.cell(row=r, column=s1_reason, value=excel_safe(res1.error))
            ws.cell(row=r, column=s1_raw, value=excel_safe(res1.raw_text))
            print(f"[SCREENING ERROR] row={r} {res1.error}", file=sys.stderr)

    if run_stage2 and stage2_cols is not None:
        if not relevant:
            status = "N/A (Stage1 Error)" if stage1_failed else "N/A (Skipped)"
            ws.cell(row=r, column=s2_status, value=status)
        else:
            res2 = await runner.generate_json(build_prompt(template2, values), args.model_stage2, validate_extraction)
            if res2.ok:
                stage2_cols.write(r, flatten_record(unwrap_extraction(res2.data) or {}))
                ws.cell(row=r, column=s2_status, value="OK")
                raw_col = stage2_cols.columns.get("2__raw_text")
                if raw_col is not None:
                    ws.cell(row=r, column=raw_col).value = None  # clear a stale error from a previous retry pass
                print(f"[EXTRACTION] row={r} OK")
            else:
                row_error = True
                ws.cell(row=r, column=s2_status, value=excel_safe(f"ERROR: {res2.error}"))
                stage2_cols.write(r, {"_raw_text": res2.raw_text})
                print(f"[EXTRACTION ERROR] row={r} {res2.error}", file=sys.stderr)

    return row_error


async def _process(job: JobConfig, args: argparse.Namespace, run_stage1: bool, run_stage2: bool) -> int:
    print(f"[PATH] 入力: {job.input_path}")
    print(f"[SHEET] {job.sheetname}")
    if job.prompt1:
        print(f"[PROMPT1] {job.prompt1} (sha1:{file_sha1(job.prompt1)[:12]})")
    if job.prompt2:
        print(f"[PROMPT2] {job.prompt2} (sha1:{file_sha1(job.prompt2)[:12]})")
    print("[AUTH] " + ("token" if job.token else "ログイン済みCopilotユーザー"))
    if args.web_search:
        print("[WARN] --web-search は Copilot SDK では無視されます", file=sys.stderr)

    template1 = read_prompt_text(job.prompt1) if job.prompt1 else ""
    template2 = read_prompt_text(job.prompt2) if job.prompt2 else ""

    suffix = Path(job.input_path).suffix.lower()
    wb = load_workbook(job.input_path, keep_vba=(suffix == ".xlsm"))
    if job.sheetname not in wb.sheetnames:
        raise Abort(f"シート '{job.sheetname}' が見つかりません。利用可能: {wb.sheetnames}")
    ws = wb[job.sheetname]

    headers = header_columns(ws)
    missing = [c for c in job.input_columns if c not in headers]
    if missing:
        raise Abort(f"シート {job.sheetname} に必要列がありません: {', '.join(missing)}")
    input_cols = [headers[c] for c in job.input_columns]
    print(f"[INPUT COLUMNS] {job.input_columns}")

    last_col = rightmost_filled_col(ws)
    print(f"右端の列番号: {last_col} （列記号: {get_column_letter(last_col) if last_col else '-'} ）")
    next_col = last_col + 1
    s1_judge = s1_reason = s1_raw = s2_status = 0
    if run_stage1:
        s1_judge, s1_reason, s1_raw = next_col, next_col + 1, next_col + 2
        ws.cell(row=1, column=s1_judge, value="1_Judgment")
        ws.cell(row=1, column=s1_reason, value="1_Reason")
        ws.cell(row=1, column=s1_raw, value="(In Case Stage1 Error: Output of LLM)")
        next_col += 3
    stage2_cols: DynamicColumns | None = None
    if run_stage2:
        s2_status = next_col
        ws.cell(row=1, column=s2_status, value="2_Status")
        stage2_cols = DynamicColumns(ws, start_col=next_col + 1, prefix="2_")

    stem = Path(job.input_path).stem
    output_path = os.path.join(job.outdir, f"{datetime.now():%Y-%m-%d_%H%M%S}_{stem}{Path(job.input_path).suffix}")
    if not _save(wb, output_path, final=True):
        return EXIT_FATAL
    print(f"{OUTPUT_MARKER} {output_path}")

    rows: list[tuple[int, list[str]]] = []
    for r in range(2, ws.max_row + 1):
        values = [cell_text(ws.cell(row=r, column=c).value) for c in input_cols]
        if any(v.strip() for v in values):
            rows.append((r, values))
    total_rows = ws.max_row - 1 if ws.max_row > 1 else 0
    print(f"[INPUT ROWS] total={total_rows} process={len(rows)} skipped_all_empty={total_rows - len(rows)}")

    row_values = dict(rows)
    pending = [r for r, _ in rows]
    total_passes = 1 + max(0, args.row_retry_passes)
    cols = (s1_judge, s1_reason, s1_raw, s2_status)
    try:
        async with CopilotRunner(
            token=job.token,
            cli_path=args.cli_path,
            cli_url=args.cli_url,
            request_timeout=args.request_timeout,
            max_retries=args.max_retries,
        ) as runner:
            for pass_no in range(1, total_passes + 1):
                if not pending:
                    break
                if pass_no > 1:
                    print(
                        f"\n[RETRY] エラーになった {len(pending)} 行を再試行します "
                        f"（{pass_no - 1}/{args.row_retry_passes} 回目）: {pending}",
                        file=sys.stderr,
                    )
                    if args.row_retry_delay > 0:
                        await asyncio.sleep(args.row_retry_delay)

                still_failing: list[int] = []
                progress = tqdm(total=len(pending), desc="rows" if pass_no == 1 else f"retry {pass_no - 1}")
                try:
                    for done, r in enumerate(pending, start=1):
                        had_error = await _process_row(
                            runner, ws, r, row_values[r], template1, template2, args,
                            run_stage1, run_stage2, cols, stage2_cols,
                        )
                        if had_error:
                            still_failing.append(r)
                        progress.update(1)
                        if done % SAVE_EVERY_ROWS == 0:
                            _save(wb, output_path)
                finally:
                    progress.close()
                    _save(wb, output_path)
                pending = still_failing
    finally:
        saved = _save(wb, output_path, final=True)
        wb.close()

    if not saved:
        return EXIT_FATAL
    if pending:
        print(f"[DONE WITH ERRORS] {len(pending)} 行でエラーが発生しました: {output_path}", file=sys.stderr)
        return EXIT_ROW_ERRORS
    print(f"[DONE] {output_path}")
    return EXIT_OK


def main(run_stage1: bool, run_stage2: bool) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    stages = " + ".join(s for s, on in (("Stage1", run_stage1), ("Stage2", run_stage2)) if on)
    args = parse_args(f"Excel（分割済み）を1本処理（{stages}）。UI/非UI両対応。")
    ui = Notifier(use_ui=not args.no_ui)

    if args.merge_only:
        return _merge_only(args, ui)

    try:
        job = _resolve_job(args, ui, run_stage1, run_stage2)
        ui.info("開始", "処理を開始します。")
        rc = asyncio.run(_process(job, args, run_stage1, run_stage2))
    except Abort as exc:
        ui.error("エラー", str(exc))
        return EXIT_FATAL
    except Exception as exc:
        traceback.print_exc()
        ui.error("エラー", f"予期せぬエラー: {type(exc).__name__}: {exc}")
        return EXIT_FATAL

    if rc == EXIT_FATAL:
        ui.error("エラー", "処理が中断されました。ログを確認してください。")
        return rc
    ui.info("完了", "すべての処理が完了しました。" if rc == EXIT_OK else "完了しましたが、一部の行でエラーが発生しました。")

    if args.merge_after:
        base_noext = Path(job.input_path).stem
        try:
            merged_path = merge_parts_into_single(
                outdir=job.outdir,
                input_base_noext=base_noext,
                sheetname=job.sheetname,
                parts=args.parts,
                run_metadata_path=args.run_metadata,
            )
            print(f"{MERGE_DONE_MARKER} {merged_path}")
            ui.info("結合完了", f"Part1〜Part{args.parts} を結合しました:\n{merged_path}")
        except Exception as exc:
            ui.error("結合エラー", str(exc))
            return EXIT_FATAL
    return rc
